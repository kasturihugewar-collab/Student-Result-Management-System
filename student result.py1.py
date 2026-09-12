from tkinter import *
from tkinter import ttk
from openpyxl import Workbook, load_workbook
import os

f="student_results.xlsx"
if not os.path.exists(f):
    wb=Workbook(); ws=wb.active
    ws.append(["Name","Roll","Class","S1","S2","S3","S4","S5","Total","%","Result"]); wb.save(f)

def save():
    wb=load_workbook(f); ws=wb.active
    m=[int(e.get()) for e in sub]
    t=sum(m); p=t/5; r="Pass" if p>=40 else "Fail"
    ws.append([name.get(),roll.get(),cls.get(),*m,t,p,r]); wb.save(f)

def get():
    wb=load_workbook(f)
    for x in wb.active.iter_rows(min_row=2,values_only=True):
        if str(x[1])==sr.get():
            out["text"]=f"{x[0]}  Total:{x[8]}  {x[9]:.1f}%  {x[10]}"; return
    out["text"]="Not Found"

def show():
    t=Toplevel(root)
    tr=ttk.Treeview(t,columns=("N","R","C","T","P","Res"),show="headings")
    for c in ("N","R","C","T","P","Res"): tr.heading(c,text=c)
    for x in load_workbook(f).active.iter_rows(min_row=2,values_only=True):
        tr.insert("",END,values=(x[0],x[1],x[2],x[8],f"{x[9]:.1f}%",x[10]))
    tr.pack(fill=BOTH,expand=True)

root=Tk(); root.title("Student Result")

Label(root,text="Name").pack(); name=Entry(root); name.pack()
Label(root,text="Roll").pack(); roll=Entry(root); roll.pack()
Label(root,text="Class").pack(); cls=Entry(root); cls.pack()

sub=[]
for i in range(5):
    Label(root,text=f"Sub {i+1}").pack()
    e=Entry(root); e.pack(); sub.append(e)

Button(root,text="Save",command=save).pack()
Label(root,text="Roll No").pack(); sr=Entry(root); sr.pack()
Button(root,text="Get Result",command=get).pack()
out=Label(root); out.pack()
Button(root,text="Show All",command=show).pack()
Button(root,text="Exit",command=root.destroy).pack()

root.mainloop()
